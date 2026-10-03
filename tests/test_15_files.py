"""Spec 15: files in (download, extract, store, use) and files out (create_document, SHARE_FILE behind approval)."""

from __future__ import annotations

import io
import json

import aiosqlite
import httpx
import pytest
from google.genai import types

from fakes import FakeSlack, agent, dm, memory_structured, mention, pdf_with
from knappy.agent.tools import ToolRegistry, current_owner
from knappy.db.repository import SqliteRepository
from knappy.files.extract import kind_of
from knappy.files.service import MAX_BYTES, DownloadError, SlackDownloader, SlackFile, TooLarge
from knappy.files.store import DocumentStore, chunks
from knappy.llm.client import to_contents
from knappy.llm.fake import FakeModel, GenerateRequest
from knappy.llm.types import Attachment, ModelTurn, UserMessage
from knappy.runtime import KnappyRuntime
from knappy.slack.egress import build_say
from knappy.slack.events import on_message
from knappy.slack.executor import SlackActionExecutor

TOKEN = "xoxb-test"
PNG = b"\x89PNG\r\n\x1a\n" + bytes(range(256)) * 4
PRICING = "Pricing: the Pro plan is $40 per seat per month, billed annually, with discounts over 100 seats."
PDF = pdf_with(
    "Q3 proposal for Acme. Scope: migrate the billing system to the new platform by September.",
    PRICING,
    "Timeline: kickoff on July 1, pilot in August, rollout complete by the end of September.",
)


class FileServer:
    """Slack's file host: serves shared files to the bot token and records every request."""

    def __init__(self) -> None:
        self.files: dict[str, bytes] = {}
        self.requests: list[httpx.Request] = []

    def share(self, name: str, mimetype: str, content: bytes, *, size: int | None = None, file_id: str | None = None) -> dict:
        file_id = file_id or f"F{len(self.files) + 1}"
        url = f"https://files.slack.com/files-pri/T-{file_id}/download/{name}"
        self.files[url] = content
        return {
            "id": file_id, "name": name, "mimetype": mimetype,
            "size": len(content) if size is None else size, "url_private_download": url,
        }

    def __call__(self, request: httpx.Request) -> httpx.Response:
        self.requests.append(request)
        if request.headers.get("authorization") != f"Bearer {TOKEN}":
            return httpx.Response(200, headers={"content-type": "text/html"}, content=b"<html>Sign in</html>")
        content = self.files.get(str(request.url))
        if content is None:
            return httpx.Response(404)
        return httpx.Response(200, headers={"content-type": "application/octet-stream"}, content=content)

    def downloader(self) -> SlackDownloader:
        return SlackDownloader(TOKEN, transport=httpx.MockTransport(self))


def runtime(repo: SqliteRepository, model, server: FileServer, client: FakeSlack | None = None) -> KnappyRuntime:
    client = client or FakeSlack()
    return KnappyRuntime(
        repo, workspace_id="T_TEST", model=model, say=build_say(client), slack=client,
        executor=SlackActionExecutor(client, DocumentStore(repo, "T_TEST")), downloader=server.downloader(),
    )


def shared(text: str, ts: str, *files: dict, user: str = "U1") -> dict:
    return {**dm(text, ts, user=user, channel=f"D{user}"), "subtype": "file_share", "files": list(files)}


def agent_requests(model: FakeModel, prefix: str) -> list[GenerateRequest]:
    return [r for r in model.requests if r.tier == "agent" and r.contents[-1].text.startswith(prefix)]


async def rows(repo: SqliteRepository, sql: str, params: tuple = ()) -> list[dict]:
    cursor = await repo.connection.execute(sql, params)
    return [dict(row) for row in await cursor.fetchall()]


async def test_file_01_pdf_is_stored_remembered_and_in_the_turn(repo: SqliteRepository) -> None:
    server = FileServer()
    model = FakeModel(agent(), structured=memory_structured())
    await runtime(repo, model, server).handle_event(shared("summarize this", "1.0", server.share("q3.pdf", "application/pdf", PDF)))

    turn = agent_requests(model, "summarize this")[0].contents[-1].text
    assert PRICING in turn and "rollout complete" in turn
    [document] = await rows(repo, "SELECT id, owner_user_id, name, text, summary FROM documents")
    assert document["owner_user_id"] == "U1" and PRICING in document["text"] and document["summary"]
    [record] = await rows(repo, "SELECT id, title, body, source FROM memory_records WHERE type = 'document'")
    assert record["title"] == "q3.pdf" and f"[[doc:{document['id']}]]" in record["body"] and record["source"] == "document"
    sources = await rows(repo, "SELECT source_type, source_id FROM memory_provenance WHERE target_id = ?", (record["id"],))
    assert ("document", document["id"]) in {(row["source_type"], row["source_id"]) for row in sources}
    assert await rows(repo, "SELECT COUNT(*) AS n FROM document_chunks WHERE document_id = ?", (document["id"],)) == [{"n": 1}]
    logged = await rows(repo, "SELECT content FROM conversation_turns WHERE role = 'user'")
    assert logged == [{"content": f"summarize this\n[Shared file q3.pdf, document_id {document['id']}]"}], "the log points, never copies"


async def test_file_02_found_from_another_conversation(repo: SqliteRepository) -> None:
    server = FileServer()
    model = FakeModel(agent(), structured=memory_structured())
    knappy = runtime(repo, model, server)
    await knappy.handle_event(shared("keep this", "1.0", server.share("q3.pdf", "application/pdf", PDF)))
    [document] = await rows(repo, "SELECT id FROM documents")

    token = current_owner.set("U1")
    try:
        assert (await knappy.tools.memory_search("q3"))[0]["type"] == "document"
        assert [item["document_id"] for item in await knappy.tools.list_files("pricing")] == [document["id"]]
        read = await knappy.tools.read_file(document["id"], query="pricing")
        whole = await knappy.tools.read_file(document["id"], max_chars=1500)
    finally:
        current_owner.reset(token)
    assert any("$40 per seat" in match for match in read["matches"])
    assert whole["truncated"] is False and PRICING in whole["text"]


async def test_read_file_query_returns_the_matching_chunk_of_a_long_document(repo: SqliteRepository) -> None:
    filler = "\n\n".join(f"Section {n}: operational notes about logistics and travel." * 4 for n in range(40))
    text = f"{filler}\n\n{PRICING}\n\n{filler}"
    server = FileServer()
    model = FakeModel(agent(), structured=memory_structured())
    knappy = runtime(repo, model, server)
    await knappy.handle_event(shared("long one", "1.0", server.share("notes.txt", "text/plain", text.encode())))
    [document] = await rows(repo, "SELECT id FROM documents")

    token = current_owner.set("U1")
    try:
        read = await knappy.tools.read_file(document["id"], query="pro plan pricing", max_chars=1500)
    finally:
        current_owner.reset(token)
    assert len(chunks(text)) > 10
    assert len(read["matches"]) == 1 and PRICING in read["matches"][0]


async def test_long_documents_go_in_as_summary_and_a_read_file_note(repo: SqliteRepository) -> None:
    server = FileServer()
    model = FakeModel(agent(), structured=memory_structured())
    big = ("All work and no play. " * 6000).encode()
    await runtime(repo, model, server).handle_event(shared("thoughts?", "1.0", server.share("big.txt", "text/plain", big)))

    turn = agent_requests(model, "thoughts?")[0].contents[-1].text
    assert len(turn) < 2000 and "read_file" in turn and "Summary:" in turn


async def test_file_03_over_the_cap_is_never_downloaded(repo: SqliteRepository) -> None:
    server = FileServer()
    client = FakeSlack()
    model = FakeModel(agent(), structured=memory_structured())
    huge = server.share("video-export.pdf", "application/pdf", b"%PDF", size=30 * 1024 * 1024)
    reply = await runtime(repo, model, server, client).handle_event(shared("read this", "1.0", huge))

    assert server.requests == [], "no download beyond the cap"
    assert "20 MB" in reply.text and "20 MB" in client.shown("100.1")["text"]
    assert await rows(repo, "SELECT id FROM documents") == []


async def test_download_stops_at_the_cap_when_slack_reports_no_size() -> None:
    server = FileServer()
    raw = server.share("dump.txt", "text/plain", b"x" * (MAX_BYTES + 1))
    with pytest.raises(TooLarge):
        await server.downloader().download(SlackFile.parse({**raw, "size": None}))


async def test_the_token_only_goes_to_slack_and_a_sign_in_page_is_not_a_file() -> None:
    server = FileServer()
    elsewhere = SlackFile("F1", "x.txt", "text/plain", 3, "https://evil.example/files/x.txt")
    with pytest.raises(DownloadError):
        await server.downloader().download(elsewhere)
    assert server.requests == []
    raw = server.share("x.txt", "text/plain", b"abc")
    with pytest.raises(DownloadError, match="files:read"):
        await SlackDownloader("xoxb-wrong", transport=httpx.MockTransport(server)).download(SlackFile.parse(raw))


async def test_file_04_image_reaches_the_model_and_only_its_description_is_kept(repo: SqliteRepository, tmp_path) -> None:
    server = FileServer()
    description = "A screenshot of a pricing table: the Pro plan is $40 per seat."

    async def respond(request: GenerateRequest) -> ModelTurn:
        if request.tier == "light":
            return ModelTurn(text=description)
        return ModelTurn(text="I see a pricing table.")

    model = FakeModel(respond, structured=memory_structured())
    database = SqliteRepository(str(tmp_path / "knappy.db"))
    await database.connect()
    await database.init_schema()
    await database.upsert_workspace("T_TEST", "Test", TOKEN)
    try:
        await runtime(database, model, server).handle_event(
            shared("what is this?", "1.0", server.share("screen.png", "image/png", PNG))
        )
        [described] = [r for r in model.requests if r.tier == "light"]
        [turn] = agent_requests(model, "what is this?")
        [document] = await rows(database, "SELECT text, summary FROM documents")
    finally:
        await database.close()

    assert described.contents[-1].attachments == (Attachment("image/png", PNG),)
    assert turn.contents[-1].attachments == (Attachment("image/png", PNG),), "the model receives an image part"
    assert document["text"] == description
    stored = (tmp_path / "knappy.db").read_bytes()
    assert PNG[8:64] not in stored, "no image bytes in the database file"


async def test_scanned_pdf_pages_are_transcribed_in_place(repo: SqliteRepository) -> None:
    server = FileServer()

    async def respond(request: GenerateRequest) -> ModelTurn:
        if request.tier == "light":
            [attachment] = request.contents[-1].attachments
            assert attachment.mime_type == "application/pdf" and attachment.data.startswith(b"%PDF")
            return ModelTurn(text="Signed by both parties on June 3.")
        return ModelTurn(text="ok")

    model = FakeModel(respond, structured=memory_structured())
    scan = pdf_with("Contract between Acme and Knappy for consulting services in 2026.", "")
    await runtime(repo, model, server).handle_event(shared("file this", "1.0", server.share("contract.pdf", "application/pdf", scan)))

    [document] = await rows(repo, "SELECT text FROM documents")
    assert document["text"] == "Contract between Acme and Knappy for consulting services in 2026.\n\nSigned by both parties on June 3."


async def test_docx_csv_and_unreadable_types(repo: SqliteRepository) -> None:
    import docx

    word = docx.Document()
    word.add_paragraph("Offsite agenda: strategy on day one.")
    table = word.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Owner", "Priya"
    buffer = io.BytesIO()
    word.save(buffer)
    sheet = "name,seats\n" + "\n".join(f"team{n},{n}" for n in range(500))
    server = FileServer()
    model = FakeModel(agent(), structured=memory_structured())
    reply = await runtime(repo, model, server).handle_event(shared(
        "file these", "1.0",
        server.share("agenda.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", buffer.getvalue()),
        server.share("seats.csv", "text/csv", sheet.encode()),
        server.share("song.mp3", "audio/mpeg", b"ID3"),
    ))

    texts = {row["name"]: row["text"] for row in await rows(repo, "SELECT name, text FROM documents")}
    assert texts["agenda.docx"] == "Offsite agenda: strategy on day one.\n\nOwner | Priya"
    assert texts["seats.csv"].splitlines()[200] == "team199,199" and "[First 200 of 500 rows.]" in texts["seats.csv"]
    assert "song.mp3" not in texts and "can't read *song.mp3* yet" in reply.text
    assert [request.url.path.rsplit("/", 1)[-1] for request in server.requests] == ["agenda.docx", "seats.csv"]


def test_kind_table() -> None:
    assert kind_of("application/pdf", "a.pdf") == "pdf"
    assert kind_of("application/octet-stream", "main.py") == "text"
    assert kind_of("", "photo.JPG") == "image"
    assert kind_of("text/x-weird", "notes") == "text"
    assert kind_of("application/zip", "a.zip") is None


async def test_sharing_the_same_file_again_reuses_the_document(repo: SqliteRepository) -> None:
    server = FileServer()
    model = FakeModel(agent(), structured=memory_structured())
    knappy = runtime(repo, model, server)
    raw = server.share("q3.pdf", "application/pdf", PDF)
    await knappy.handle_event(shared("summarize this", "1.0", raw))
    await knappy.handle_event(shared("and again", "2.0", raw))

    assert len(server.requests) == 1
    assert len(await rows(repo, "SELECT id FROM documents")) == 1
    assert PRICING in agent_requests(model, "and again")[0].contents[-1].text


async def test_file_05_create_document_uploads_to_the_requesters_dm(repo: SqliteRepository) -> None:
    plan = "# Project plan\n- Scope\n- Milestones"
    client = FakeSlack()
    model = FakeModel(agent({"a one-page project plan": ("create_document", {"title": "Project plan", "content_markdown": plan})}),
                      structured=memory_structured())
    reply = await runtime(repo, model, FileServer(), client).handle_event(mention("a one-page project plan please", "1.0"))

    assert [(upload["channel"], upload["filename"], upload["content"]) for upload in client.uploads] == [
        ("DU1", "project-plan.md", plan)
    ], "a mention in a channel still delivers to the requester's own DM"
    assert reply.blocks is None and await rows(repo, "SELECT id FROM action_drafts") == []
    result = json.loads(reply.text)[0]
    assert result["permalink"] == "https://slack.test/files/FUP1"
    assert await rows(repo, "SELECT name, owner_user_id FROM documents") == [{"name": "project-plan.md", "owner_user_id": "U1"}]


async def test_file_06_share_file_waits_for_approval(repo: SqliteRepository) -> None:
    plan = "# Launch plan\n- Owners\n- Dates"
    client = FakeSlack()

    stored_id: list[str] = []

    def share_args(request: GenerateRequest) -> dict:
        return {"action_type": "SHARE_FILE", "recipient": "Alex", "recipient_identifier": "UALEX",
                "summary": "Send Alex the launch plan", "staged_content": "Here's the launch plan.", "document_id": stored_id[0]}

    model = FakeModel(agent({
        "write the launch plan": ("create_document", {"title": "Launch plan", "content_markdown": plan}),
        "send that plan to alex": ("stage_outbound_action", share_args),
    }), structured=memory_structured())
    knappy = runtime(repo, model, FileServer(), client)
    await knappy.handle_event(dm("write the launch plan", "1.0", channel="DU1"))
    stored_id.append((await rows(repo, "SELECT id FROM documents"))[0]["id"])
    card = await knappy.handle_event(dm("send that plan to Alex", "2.0", channel="DU1"))

    assert card.draft_id and "btn_approve_action" in json.dumps(card.blocks)
    assert "*File:* launch-plan.md" in json.dumps(card.blocks)
    before = [upload["channel"] for upload in client.uploads]
    assert before == ["DU1"] and not [post for post in client.posts if post["channel"] in ("UALEX", "DUALEX")], "nothing before approval"
    assert (await knappy.gateway.approve(card.draft_id, "U2")).executed is False
    assert [upload["channel"] for upload in client.uploads] == ["DU1"], "only the owner can approve"

    approved = await knappy.gateway.approve(card.draft_id, "U1")
    again = await knappy.gateway.approve(card.draft_id, "U1")
    sent = [upload for upload in client.uploads if upload["channel"] == "DUALEX"]
    assert approved.executed and not again.executed
    assert [(upload["content"], upload["initial_comment"]) for upload in sent] == [(plan, "Here's the launch plan.")]


async def test_file_07_owners_cannot_reach_each_others_documents(repo: SqliteRepository) -> None:
    server = FileServer()
    model = FakeModel(agent(), structured=memory_structured())
    knappy = runtime(repo, model, server)
    await knappy.handle_event(shared("mine", "1.0", server.share("q3.pdf", "application/pdf", PDF), user="U1"))
    [document] = await rows(repo, "SELECT id FROM documents")

    token = current_owner.set("U2")
    try:
        assert await knappy.tools.read_file(document["id"]) == {"error": f"No document with id {document['id']}"}
        assert await knappy.tools.read_file(document["id"], query="pricing") == {"error": f"No document with id {document['id']}"}
        assert await knappy.tools.list_files() == []
        assert await knappy.tools.list_files("pricing") == []
        assert await knappy.tools.memory_search("q3") == []
    finally:
        current_owner.reset(token)
    assert await knappy.documents.best_chunks("U2", document["id"], "pricing", 20_000) == []


async def test_share_file_refuses_another_owners_document(repo: SqliteRepository) -> None:
    server = FileServer()
    model = FakeModel(agent(), structured=memory_structured())
    knappy = runtime(repo, model, server)
    await knappy.handle_event(shared("mine", "1.0", server.share("q3.pdf", "application/pdf", PDF), user="U1"))
    [document] = await rows(repo, "SELECT id FROM documents")

    model2 = FakeModel(agent({"send": ("stage_outbound_action", {
        "action_type": "SHARE_FILE", "recipient": "Alex", "recipient_identifier": "UALEX", "summary": "s",
        "staged_content": "here", "document_id": document["id"]})}), structured=memory_structured())
    reply = await runtime(repo, model2, server).handle_event(dm("send it to Alex", "2.0", user="U2", channel="DU2"))

    assert reply.draft_id is None and "No document with id" in reply.text
    assert await rows(repo, "SELECT id FROM action_drafts") == []


async def test_forgetting_a_document_deletes_it_and_what_it_fed(repo: SqliteRepository) -> None:
    server = FileServer()
    model = FakeModel(agent(), structured=memory_structured())
    knappy = runtime(repo, model, server)
    await knappy.handle_event(shared("keep", "1.0", server.share("q3.pdf", "application/pdf", PDF)))
    [document] = await rows(repo, "SELECT id FROM documents")
    [record] = await rows(repo, "SELECT id FROM memory_records WHERE type = 'document'")
    store = knappy.store
    async with repo.transaction():
        await store.create_record(
            "U1", record_id="fact:acme-price", type="fact", title="Acme price", body="- Pro plan is $40 per seat",
            source="reconciler", now=store.clock(), sources=[("record", record["id"])],
        )

    result = await knappy.memory_engine.forget("U1", record["id"])

    assert result["forgotten"] == ["q3.pdf"] and "fact:acme-price" in result["also_forgotten"]
    assert await rows(repo, "SELECT id FROM documents") == []
    assert await rows(repo, "SELECT seq FROM document_chunks") == []
    token = current_owner.set("U1")
    try:
        assert "error" in await knappy.tools.read_file(document["id"])
    finally:
        current_owner.reset(token)


async def test_old_databases_gain_share_file(tmp_path) -> None:
    path = str(tmp_path / "old.db")
    async with aiosqlite.connect(path) as old:
        await old.executescript(
            """
            CREATE TABLE workspaces (id TEXT PRIMARY KEY, team_name TEXT NOT NULL, bot_token TEXT NOT NULL,
                installed_at DATETIME DEFAULT CURRENT_TIMESTAMP);
            INSERT INTO workspaces (id, team_name, bot_token) VALUES ('T_TEST', 't', 'x');
            CREATE TABLE action_drafts (
                id TEXT PRIMARY KEY,
                workspace_id TEXT NOT NULL REFERENCES workspaces(id) ON DELETE CASCADE,
                user_id TEXT NOT NULL, channel_id TEXT NOT NULL, thread_ts TEXT,
                action_type TEXT NOT NULL CHECK(action_type IN ('SEND_SLACK_DM', 'GMAIL_DRAFT', 'CALENDAR_INVITE', 'POST_CHANNEL')),
                payload TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'PENDING' CHECK(status IN ('PENDING', 'APPROVED', 'CANCELLED', 'EXPIRED', 'FAILED')),
                expires_at DATETIME, created_at DATETIME DEFAULT CURRENT_TIMESTAMP, executed_at DATETIME
            );
            INSERT INTO action_drafts (id, workspace_id, user_id, channel_id, action_type, payload)
            VALUES ('d1', 'T_TEST', 'U1', 'D1', 'SEND_SLACK_DM', '{}');
            """
        )
        await old.commit()
    for _ in range(2):
        database = SqliteRepository(path)
        await database.connect()
        await database.init_schema()
        await database.close()
    database = SqliteRepository(path)
    await database.connect()
    try:
        draft = await database.create_draft(
            workspace_id="T_TEST", user_id="U1", channel_id="D1", action_type="SHARE_FILE", payload={}
        )
        kept = await rows(database, "SELECT id, action_type FROM action_drafts ORDER BY id")
        indexes = await rows(database, "SELECT name FROM sqlite_master WHERE type = 'index' AND tbl_name = 'action_drafts'")
    finally:
        await database.close()
    assert {row["id"]: row["action_type"] for row in kept} == {"d1": "SEND_SLACK_DM", draft: "SHARE_FILE"}
    assert {"name": "idx_action_drafts_pending"} in indexes


def test_attachments_become_inline_parts() -> None:
    [content] = to_contents([UserMessage("what is this?", (Attachment("image/png", PNG),))])
    assert content.parts[0].text == "what is this?"
    assert content.parts[1].inline_data == types.Blob(data=PNG, mime_type="image/png")


async def test_file_share_messages_reach_the_agent_and_edits_do_not() -> None:
    seen: list[dict] = []

    async def processor(event: dict) -> None:
        seen.append(event)

    async def ack() -> None:
        return None

    await on_message({"subtype": "file_share", "channel_type": "im", "user": "U1", "ts": "1.0", "files": []}, ack, processor=processor)
    await on_message({"subtype": "message_changed", "channel_type": "im", "ts": "2.0"}, ack, processor=processor)
    assert [event["ts"] for event in seen] == ["1.0"]


async def test_tools_without_files_say_so(repo: SqliteRepository) -> None:
    registry = ToolRegistry(repo, "T_TEST")
    assert "error" in await registry.read_file("doc_x")
    assert "error" in await registry.list_files()
