"""What a shared file says, by type (Spec 15 §2.2). Pure and synchronous; callers run it in a thread.

Extraction returns pieces in reading order. Text is read here; an Attachment is a piece only the model can read
(an image, or a PDF page with no text layer), and the caller swaps it for the model's description.
"""

from __future__ import annotations

import io
import mimetypes
from pathlib import PurePosixPath
from typing import Literal

from knappy.llm.types import Attachment

Kind = Literal["pdf", "docx", "text", "csv", "image"]
Piece = str | Attachment

SCANNED_PAGE_CHARS = 50
MAX_SCANNED_PAGES = 20
CSV_ROWS = 200
MAX_TEXT_CHARS = 500_000

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
IMAGE_MIMES = frozenset({"image/png", "image/jpeg", "image/gif", "image/webp"})
MIME_KINDS: dict[str, Kind] = {
    "application/pdf": "pdf",
    DOCX_MIME: "docx",
    "text/csv": "csv",
    "application/json": "text",
    "application/xml": "text",
    "application/x-yaml": "text",
    "application/javascript": "text",
    "application/x-sh": "text",
    **{mime: "image" for mime in IMAGE_MIMES},
}
EXTENSION_KINDS: dict[str, Kind] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".csv": "csv",
    **{ext: "image" for ext in (".png", ".jpg", ".jpeg", ".gif", ".webp")},
    **{
        ext: "text"
        for ext in (
            ".txt", ".md", ".markdown", ".json", ".yaml", ".yml", ".toml", ".ini", ".xml", ".html", ".css", ".sql",
            ".py", ".js", ".jsx", ".ts", ".tsx", ".go", ".rs", ".java", ".kt", ".rb", ".php", ".c", ".h", ".cpp",
            ".cs", ".swift", ".sh", ".log",
        )
    },
}


class UnreadableError(Exception):
    pass


def kind_of(mimetype: str, name: str) -> Kind | None:
    """The extraction route for a file, or None when it isn't readable yet."""
    mime = mimetype.split(";", 1)[0].strip().lower()
    if mime in MIME_KINDS:
        return MIME_KINDS[mime]
    by_extension = EXTENSION_KINDS.get(PurePosixPath(name.lower()).suffix)
    if by_extension is not None:
        return by_extension
    return "text" if mime.startswith("text/") else None


def image_mime(mimetype: str, name: str) -> str:
    mime = mimetype.split(";", 1)[0].strip().lower()
    return mime if mime in IMAGE_MIMES else mimetypes.guess_type(name)[0] or "image/png"


def extract(kind: Kind, data: bytes, mime: str) -> list[Piece]:
    if kind == "image":
        return [Attachment(mime, data)]
    if kind == "pdf":
        return pdf_pieces(data)
    if kind == "docx":
        return [_docx_text(data)]
    if kind == "csv":
        return [_csv_text(data)]
    return [_decode(data)]


def pdf_page_texts(data: bytes) -> list[str]:
    return [text for text, _page in _pdf_pages(data)]


def pdf_pieces(data: bytes) -> list[Piece]:
    """Page text, or the page itself as a one-page PDF when it has under 50 characters of text (a scan)."""
    from pypdf import PdfWriter

    pieces: list[Piece] = []
    scanned = 0
    for number, (text, page) in enumerate(_pdf_pages(data), 1):
        if len(text.strip()) >= SCANNED_PAGE_CHARS:
            pieces.append(text)
        elif scanned < MAX_SCANNED_PAGES:
            scanned += 1
            writer = PdfWriter()
            writer.add_page(page)
            buffer = io.BytesIO()
            writer.write(buffer)
            pieces.append(Attachment("application/pdf", buffer.getvalue()))
        else:
            pieces.append(f"[Page {number} is scanned and was not transcribed.]")
    return pieces


def _pdf_pages(data: bytes) -> list[tuple[str, object]]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        return [((page.extract_text() or "").strip(), page) for page in reader.pages]
    except (PdfReadError, ValueError, KeyError) as exc:
        raise UnreadableError(f"Could not read the PDF: {exc}") from None


def _docx_text(data: bytes) -> str:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, ValueError, KeyError) as exc:
        raise UnreadableError(f"Could not read the Word document: {exc}") from None
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        parts.append("\n".join(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows))
    return "\n\n".join(parts)


def _csv_text(data: bytes) -> str:
    lines = _decode(data).splitlines()
    kept = "\n".join(lines[: CSV_ROWS + 1])
    if len(lines) > CSV_ROWS + 1:
        kept += f"\n[First {CSV_ROWS} of {len(lines) - 1} rows.]"
    return kept


def _decode(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")
